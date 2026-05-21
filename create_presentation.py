from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import copy

# Color palette
BG_DARK = RGBColor(0x0A, 0x0A, 0x14)        # Near black
BG_CARD = RGBColor(0x12, 0x12, 0x24)        # Dark card
ACCENT = RGBColor(0x00, 0xD4, 0xAA)         # Teal/green
ACCENT2 = RGBColor(0x7B, 0x5E, 0xFF)        # Purple
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xB0, 0xB8, 0xCC)
MID_GRAY = RGBColor(0x5A, 0x62, 0x7A)

W = Inches(13.33)
H = Inches(7.5)


def new_prs():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs


def blank_layout(prs):
    return prs.slide_layouts[6]  # completely blank


def add_rect(slide, x, y, w, h, fill_color, border_color=None, border_width=None):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width or Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_text(slide, text, x, y, w, h, font_size, bold=False, color=WHITE,
             align=PP_ALIGN.LEFT, italic=False, wrap=True):
    txb = slide.shapes.add_textbox(x, y, w, h)
    txb.word_wrap = wrap
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return txb


def add_bullet_text(slide, items, x, y, w, h, font_size=14, color=LIGHT_GRAY,
                    bullet_color=ACCENT, title=None, title_size=18):
    txb = slide.shapes.add_textbox(x, y, w, h)
    txb.word_wrap = True
    tf = txb.text_frame
    tf.word_wrap = True

    first = True
    if title:
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = title
        run.font.size = Pt(title_size)
        run.font.bold = True
        run.font.color.rgb = WHITE
        run.font.name = "Calibri"
        first = False

    for item in items:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(4)
        run = p.add_run()
        run.text = f"• {item}"
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
        run.font.name = "Calibri"


def set_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_accent_line(slide, x, y, w, color=ACCENT):
    add_rect(slide, x, y, w, Pt(3), color)


def add_tag(slide, text, x, y, color=ACCENT):
    w = Inches(1.8)
    h = Inches(0.32)
    add_rect(slide, x, y, w, h, color)
    add_text(slide, text, x, y, w, h, 10, bold=True, color=BG_DARK, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────────
# BUILD PRESENTATION
# ─────────────────────────────────────────────────────────────
prs = new_prs()
blank = blank_layout(prs)

# ── SLIDE 1: COVER ───────────────────────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

# Top label
add_tag(slide, "US Product Offerings", Inches(0.6), Inches(0.5))

# Big headline
add_text(slide, "The Modern Bank,", Inches(0.6), Inches(1.5), Inches(7), Inches(1.1),
         52, bold=True, color=WHITE)
add_text(slide, "built for the", Inches(0.6), Inches(2.5), Inches(6), Inches(0.9),
         52, bold=True, color=WHITE)
add_text(slide, "Modern Businesses.", Inches(0.6), Inches(3.3), Inches(8), Inches(1.0),
         52, bold=True, color=ACCENT)

# Bottom domain
add_text(slide, "STARKBANK.COM", Inches(0.6), Inches(6.7), Inches(4), Inches(0.5),
         11, bold=True, color=MID_GRAY, align=PP_ALIGN.LEFT)

# Decorative accent block (right side)
add_rect(slide, Inches(10.5), Inches(0), Inches(2.83), Inches(7.5), RGBColor(0x12, 0x12, 0x28))
add_rect(slide, Inches(12.5), Inches(2), Inches(0.83), Inches(3.5), ACCENT)

# ── SLIDE 2: ABOUT ───────────────────────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "About Stark Bank", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3))

add_text(slide,
         "Stark Bank is the bank that transforms\ncompanies through technology, offering\ninnovative banking services for the B2B market.",
         Inches(0.6), Inches(1.3), Inches(7.5), Inches(2.5),
         28, bold=True, color=WHITE)

# Cards: product categories
categories = [
    ("Business Accounts", ACCENT),
    ("Accounts Payables", ACCENT2),
    ("Certificates of Deposits", ACCENT),
    ("Accounts Receivables", ACCENT2),
    ("Corporate Credit Card", ACCENT),
    ("Smart Lines of Credit", ACCENT2),
]
cols = 3
card_w = Inches(3.8)
card_h = Inches(0.9)
gap_x = Inches(0.25)
gap_y = Inches(0.2)
start_x = Inches(0.55)
start_y = Inches(4.3)

for i, (label, color) in enumerate(categories):
    col = i % cols
    row = i // cols
    x = start_x + col * (card_w + gap_x)
    y = start_y + row * (card_h + gap_y)
    add_rect(slide, x, y, card_w, card_h, BG_CARD)
    add_rect(slide, x, y, Inches(0.06), card_h, color)
    add_text(slide, label, x + Inches(0.15), y, card_w - Inches(0.2), card_h,
             15, bold=True, color=WHITE)

# ── SLIDE 3: TOTAL CONTROL ───────────────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Operations", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3))

add_text(slide, "Total control\nof your operations.", Inches(0.6), Inches(1.2),
         Inches(6.5), Inches(1.8), 36, bold=True, color=WHITE)
add_text(slide, "Create multiple accounts & user permissions.",
         Inches(0.6), Inches(2.9), Inches(7), Inches(0.6), 18, color=LIGHT_GRAY)

bullets = [
    "Sub-accounts to track spend across teams, clients, or projects",
    "Create custom rules to approve payments based on your governance policy",
    "Role-based access for teams",
    "Full audit logs on every action",
]
add_bullet_text(slide, bullets, Inches(0.6), Inches(3.7), Inches(7.5), Inches(3.2),
                font_size=16, color=LIGHT_GRAY)

# Right visual block
add_rect(slide, Inches(9.5), Inches(1.5), Inches(3.5), Inches(5.5), BG_CARD)
add_text(slide, "[ Dashboard Preview ]", Inches(9.5), Inches(3.5), Inches(3.5), Inches(1),
         14, color=MID_GRAY, align=PP_ALIGN.CENTER)

# ── SLIDE 4: RECEIVABLES ─────────────────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Receivables", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3))

add_text(slide, "Receivables\nall in one place.", Inches(0.6), Inches(1.2),
         Inches(6.5), Inches(1.8), 36, bold=True, color=WHITE)
add_text(slide,
         "From sales to collections, every dollar flows\nseamlessly — tracked, reconciled, and verified\nthe moment it moves.",
         Inches(0.6), Inches(2.9), Inches(6), Inches(1.5), 18, color=LIGHT_GRAY)

bullets = [
    "ACH, Wire, and FedNow for instant or next-day settlement",
    "Invoices with automatic reconciliation",
    "Integrated card checkout",
    "Custom payment links to collect instantly",
]
add_bullet_text(slide, bullets, Inches(0.6), Inches(4.5), Inches(6.5), Inches(2.8),
                font_size=16, color=LIGHT_GRAY)

add_rect(slide, Inches(9.5), Inches(1.5), Inches(3.5), Inches(5.5), BG_CARD)
add_text(slide, "[ Receivables Flow ]", Inches(9.5), Inches(3.5), Inches(3.5), Inches(1),
         14, color=MID_GRAY, align=PP_ALIGN.CENTER)

# ── SLIDE 5: INVOICE RECONCILIATION ──────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Accounts Receivable", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3.5))

add_text(slide, "Invoice Reconciliation.\nAutomated Matching.", Inches(0.6), Inches(1.2),
         Inches(7), Inches(1.8), 36, bold=True, color=WHITE)

bullets = [
    "Every incoming payment (ACH, FedNow, Wire) is automatically matched to the correct invoice using payment metadata and bank account identifiers",
    "Eliminates manual searches through statements",
    "Collection Dashboards to track overdue invoices",
    "Track and find every payment on your statement using tags",
]
add_bullet_text(slide, bullets, Inches(0.6), Inches(3.2), Inches(7.5), Inches(4),
                font_size=16, color=LIGHT_GRAY)

add_rect(slide, Inches(9.5), Inches(1.5), Inches(3.5), Inches(5.5), BG_CARD)
add_text(slide, "[ Reconciliation View ]", Inches(9.5), Inches(3.5), Inches(3.5), Inches(1),
         14, color=MID_GRAY, align=PP_ALIGN.CENTER)

# ── SLIDE 6: PAY ANYONE ──────────────────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Payments", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3))

add_text(slide, "Pay Anyone, Anytime.\nAll in one place.", Inches(0.6), Inches(1.2),
         Inches(7), Inches(1.8), 36, bold=True, color=WHITE)
add_text(slide,
         "Whether it's payroll, vendor payments, taxes, utilities, or refunds,\nevery transaction is tracked, approved, and settled in real time.",
         Inches(0.6), Inches(2.9), Inches(7.5), Inches(1.1), 16, color=LIGHT_GRAY)

bullets = [
    "ACH, Wire and FedNow for vendor payments and operating expenses",
    "Bill pay for vendors and utilities",
    "Bulk disbursements for payroll, contractors, or mass payouts",
    "Review, approve and schedule payments",
    "Real-time ledger to track all your transactions",
]
add_bullet_text(slide, bullets, Inches(0.6), Inches(4.1), Inches(7.5), Inches(3),
                font_size=15, color=LIGHT_GRAY)

add_rect(slide, Inches(9.5), Inches(1.5), Inches(3.5), Inches(5.5), BG_CARD)
add_text(slide, "[ Payment Dashboard ]", Inches(9.5), Inches(3.5), Inches(3.5), Inches(1),
         14, color=MID_GRAY, align=PP_ALIGN.CENTER)

# ── SLIDE 7: TAGS ────────────────────────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Payments", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3))

add_text(slide, "Tag it your\nown way.", Inches(0.6), Inches(1.2), Inches(6), Inches(1.8),
         40, bold=True, color=WHITE)
add_text(slide, "We've upgraded the way you find your payments.",
         Inches(0.6), Inches(3.0), Inches(7), Inches(0.6), 18, color=LIGHT_GRAY)
add_text(slide,
         "Add custom tags to every transaction — making it easy to organize,\nsearch, and reconcile directly from your statement.",
         Inches(0.6), Inches(3.8), Inches(7.5), Inches(1.2), 16, color=LIGHT_GRAY)

# Sample tag chips
tags = ["#payroll", "#vendor-01", "#marketing", "#Q2-2025", "#utilities"]
tag_x = Inches(0.6)
tag_y = Inches(5.5)
for tag in tags:
    add_rect(slide, tag_x, tag_y, Inches(1.5), Inches(0.4), BG_CARD,
             border_color=ACCENT, border_width=Pt(1))
    add_text(slide, tag, tag_x, tag_y, Inches(1.5), Inches(0.4),
             11, color=ACCENT, align=PP_ALIGN.CENTER)
    tag_x += Inches(1.65)

# ── SLIDE 8: PAYMENT CONTROLS ────────────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Payments", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3))

add_text(slide, "Control over\nyour payments.", Inches(0.6), Inches(1.2), Inches(6), Inches(1.8),
         40, bold=True, color=WHITE)
add_text(slide, "More autonomy and control for your company.",
         Inches(0.6), Inches(3.0), Inches(7), Inches(0.6), 18, color=LIGHT_GRAY)
add_text(slide,
         "Create flexible approval rules according to your business needs.\nChoose individuals or groups responsible for approving payments\nbased on specific amount.",
         Inches(0.6), Inches(3.7), Inches(7.5), Inches(1.5), 16, color=LIGHT_GRAY)

add_rect(slide, Inches(9.5), Inches(1.5), Inches(3.5), Inches(5.5), BG_CARD)
add_text(slide, "[ Approval Rules ]", Inches(9.5), Inches(3.5), Inches(3.5), Inches(1),
         14, color=MID_GRAY, align=PP_ALIGN.CENTER)

# ── SLIDE 9: MAILBOX BILL PAY ────────────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Payments", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3))

# Two feature cards side by side
features = [
    {
        "title": "MailBox\nBill Pay",
        "subtitle": "Making payments simpler and faster",
        "body": "Invoices and bills received by email are automatically turned into ready-to-approve payments inside the internet banking — no manual input required.",
        "color": ACCENT,
    },
    {
        "title": "Vision Capture\nDocuments",
        "subtitle": "Drag. Drop. Done.",
        "body": "Our smart banking tool reads PDFs or images, auto-fills payment details, and routes them for approval — no manual input needed.",
        "color": ACCENT2,
    },
]

card_w = Inches(5.8)
card_h = Inches(5.5)
for i, f in enumerate(features):
    x = Inches(0.5) + i * (card_w + Inches(0.4))
    y = Inches(1.3)
    add_rect(slide, x, y, card_w, card_h, BG_CARD)
    add_rect(slide, x, y, card_w, Inches(0.06), f["color"])
    add_text(slide, f["title"], x + Inches(0.3), y + Inches(0.25), card_w - Inches(0.6),
             Inches(1.0), 26, bold=True, color=WHITE)
    add_text(slide, f["subtitle"], x + Inches(0.3), y + Inches(1.3), card_w - Inches(0.6),
             Inches(0.5), 14, bold=True, color=f["color"])
    add_text(slide, f["body"], x + Inches(0.3), y + Inches(1.9), card_w - Inches(0.6),
             Inches(3.2), 14, color=LIGHT_GRAY)

# ── SLIDE 10: WEB BANKING WITH SPREADSHEETS ──────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Payments", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3))

add_text(slide, "Web Banking\nwith Spreadsheets.", Inches(0.6), Inches(1.2),
         Inches(7), Inches(1.8), 36, bold=True, color=WHITE)
add_text(slide,
         "Control your account inside Excel or Google Sheets.\nYour accounting team deserves the convenience of banking\nwith sheets through an intuitive interface.",
         Inches(0.6), Inches(3.0), Inches(6.5), Inches(1.5), 16, color=LIGHT_GRAY)

bullets = [
    "Get balance and statements",
    "Request payments",
    "Issue and manage invoices",
    "Request physical credit card",
    "Real-time integration",
]
add_bullet_text(slide, bullets, Inches(0.6), Inches(4.6), Inches(6), Inches(2.5),
                font_size=15, color=LIGHT_GRAY)

add_rect(slide, Inches(9.5), Inches(1.5), Inches(3.5), Inches(5.5), BG_CARD)
add_text(slide, "[ Spreadsheet Integration ]", Inches(9.5), Inches(3.5), Inches(3.5), Inches(1),
         13, color=MID_GRAY, align=PP_ALIGN.CENTER)

# ── SLIDE 11: CORPORATE CARD COVER ───────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_tag(slide, "Corporate Card", Inches(0.6), Inches(0.5), color=ACCENT2)

add_text(slide, "Corporate Cards that\nthink like a CFO.", Inches(0.6), Inches(1.5),
         Inches(8), Inches(2.2), 44, bold=True, color=WHITE)
add_text(slide, "Turn everyday spend into a business advantage.",
         Inches(0.6), Inches(3.8), Inches(7), Inches(0.6), 20, color=LIGHT_GRAY)

add_rect(slide, Inches(10.5), Inches(0), Inches(2.83), Inches(7.5), RGBColor(0x12, 0x12, 0x28))
add_rect(slide, Inches(12.5), Inches(2), Inches(0.83), Inches(3.5), ACCENT2)

# ── SLIDE 12: CORPORATE CARD FEATURES ────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Corporate Card", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT2)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3), color=ACCENT2)

add_text(slide, "Turn everyday spend into\na business advantage.", Inches(0.6), Inches(1.2),
         Inches(7), Inches(1.6), 32, bold=True, color=WHITE)

bullets = [
    "Create virtual cards per vendor, subscription, or employee",
    "Set limits by amount, merchant, geography, time of day, and team",
    "See spend in real time and flag anything off-policy",
    "Categorize purchases to track expenses",
    "Sync with your ERP and attach receipts for audits",
    "Just-in-time funding and full ledger sync to reduce risk",
]
add_bullet_text(slide, bullets, Inches(0.6), Inches(3.0), Inches(7.5), Inches(4.2),
                font_size=15, color=LIGHT_GRAY)

add_rect(slide, Inches(9.5), Inches(1.5), Inches(3.5), Inches(5.5), BG_CARD)
add_text(slide, "[ Card Controls ]", Inches(9.5), Inches(3.5), Inches(3.5), Inches(1),
         14, color=MID_GRAY, align=PP_ALIGN.CENTER)

# ── SLIDE 13: CREDIT LINE ────────────────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Smart Lines of Credit", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(4))

add_text(slide, "Funded Your Way.", Inches(0.6), Inches(1.2), Inches(8), Inches(1.0),
         40, bold=True, color=WHITE)
add_text(slide, "Access revolving capital without giving up equity.",
         Inches(0.6), Inches(2.3), Inches(7.5), Inches(0.6), 20, color=LIGHT_GRAY)
add_text(slide,
         "Skip the dilution of venture funding or the rigid terms of short-term loans —\nStark's credit lines grow with your cash flow, not your cap table.",
         Inches(0.6), Inches(3.0), Inches(8), Inches(1.0), 16, color=LIGHT_GRAY)

steps = [
    ("1", "Start fully secured", "Backed by your deposits and investments"),
    ("2", "Graduate to unsecured", "Based on real-time behavior and financial performance"),
    ("3", "Manage seamlessly", "Through internet banking or our mobile app"),
]
step_y = Inches(4.3)
for num, title, body in steps:
    add_rect(slide, Inches(0.6), step_y, Inches(0.5), Inches(0.5), ACCENT)
    add_text(slide, num, Inches(0.6), step_y, Inches(0.5), Inches(0.5),
             16, bold=True, color=BG_DARK, align=PP_ALIGN.CENTER)
    add_text(slide, title, Inches(1.3), step_y, Inches(5), Inches(0.4),
             15, bold=True, color=WHITE)
    add_text(slide, body, Inches(1.3), step_y + Inches(0.35), Inches(6.5), Inches(0.4),
             13, color=LIGHT_GRAY)
    step_y += Inches(0.95)

# ── SLIDE 14: PROGRAMMABLE SPENDING ──────────────────────────
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Corporate Card", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT2)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3), color=ACCENT2)

add_text(slide, "Programmable spending\nreduces waste.", Inches(0.6), Inches(1.2),
         Inches(7), Inches(1.6), 34, bold=True, color=WHITE)
add_text(slide, "Define exactly where, when, and how money flows.",
         Inches(0.6), Inches(2.9), Inches(7), Inches(0.5), 18, color=LIGHT_GRAY)

controls = [
    ("Limit", "Set spending caps by amount and time period", ACCENT),
    ("Date & Time", "Control when the card can be used — by day and hour", ACCENT2),
    ("Category", "Define allowed merchant categories (MCC), physical or virtual", ACCENT),
    ("Country", "Restrict card usage to specific countries", ACCENT2),
]
ctrl_y = Inches(3.7)
ctrl_w = Inches(5.8)
ctrl_h = Inches(0.75)
for label, desc, color in controls:
    add_rect(slide, Inches(0.6), ctrl_y, ctrl_w, ctrl_h, BG_CARD)
    add_rect(slide, Inches(0.6), ctrl_y, Inches(0.07), ctrl_h, color)
    add_text(slide, label, Inches(0.85), ctrl_y + Inches(0.05), Inches(1.5), ctrl_h,
             13, bold=True, color=WHITE)
    add_text(slide, desc, Inches(0.85), ctrl_y + Inches(0.3), ctrl_w - Inches(0.4), ctrl_h,
             12, color=LIGHT_GRAY)
    ctrl_y += Inches(0.85)

add_rect(slide, Inches(9.5), Inches(1.5), Inches(3.5), Inches(5.5), BG_CARD)
add_text(slide, "[ Spending Controls ]", Inches(9.5), Inches(3.5), Inches(3.5), Inches(1),
         14, color=MID_GRAY, align=PP_ALIGN.CENTER)

# ── SLIDE 15: CARD EXTRAS (Tokenization, Digital Wallets, Receipts) ──
slide = prs.slides.add_slide(blank)
set_bg(slide)

add_text(slide, "Corporate Card", Inches(0.6), Inches(0.5), Inches(8), Inches(0.5),
         13, bold=True, color=ACCENT2)
add_accent_line(slide, Inches(0.6), Inches(1.05), Inches(3), color=ACCENT2)

add_text(slide, "More than a card.", Inches(0.6), Inches(1.2), Inches(7), Inches(0.8),
         36, bold=True, color=WHITE)

extras = [
    {
        "title": "Corporate Card\nDashboards",
        "body": "Take full control of your company's spend. Rich dashboards to analyze transactions per cardholder, purchase category and cost center.",
        "color": ACCENT,
    },
    {
        "title": "Tokenization\nper Merchant",
        "body": "Reduce fraud by linking a single virtual card to a specific merchant — unique tokenization per vendor.",
        "color": ACCENT2,
    },
    {
        "title": "Digital Wallets &\nReceipts",
        "body": "Add cards to Apple Pay or Google Pay. Attach receipts directly in the mobile app — no more chasing documents for audits.",
        "color": RGBColor(0xFF, 0x8C, 0x00),
    },
]

card_w = Inches(3.8)
card_h = Inches(4.8)
gap = Inches(0.3)
start_x = Inches(0.5)
card_y = Inches(2.1)

for i, e in enumerate(extras):
    x = start_x + i * (card_w + gap)
    add_rect(slide, x, card_y, card_w, card_h, BG_CARD)
    add_rect(slide, x, card_y, card_w, Inches(0.06), e["color"])
    add_text(slide, e["title"], x + Inches(0.25), card_y + Inches(0.2),
             card_w - Inches(0.4), Inches(0.9), 20, bold=True, color=WHITE)
    add_text(slide, e["body"], x + Inches(0.25), card_y + Inches(1.2),
             card_w - Inches(0.4), Inches(3.3), 14, color=LIGHT_GRAY)

# ── SAVE ─────────────────────────────────────────────────────
out = "/home/user/NEW/Stark_Bank_US_Product_Offering_Template.pptx"
prs.save(out)
print(f"Saved: {out}")
